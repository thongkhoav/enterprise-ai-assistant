from __future__ import annotations

import re
from dataclasses import dataclass, field
from hashlib import sha1
from io import BytesIO
from pathlib import Path

from fastapi import UploadFile
from pypdf import PdfReader

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".md", ".txt"}

# A line that may be a page number footer or header: "Page 2", "Page 2 of 5", "2 of 5", "2 / 5".
_PAGE_NUMBER_LINE = re.compile(
    r"^[ \t]*(page[ \t]+)?(\d+)(?:[ \t]*(?:of|/)[ \t]*(\d+))?[ \t]*$",
    re.IGNORECASE | re.MULTILINE,
)


@dataclass(frozen=True)
class LoadedDocument:
    document_id: str
    source: str
    # One entry per PDF page; a single entry for every other file type.
    pages: list[str]
    metadata: dict = field(default_factory=dict)
    paged: bool = False

    @property
    def text(self) -> str:
        return "\n\n".join(self.pages)


async def load_upload(file: UploadFile, *, source: str | None = None) -> LoadedDocument:
    """Read an uploaded file, extract its text, and return a LoadedDocument.

    The document_id is a content hash, so uploading the same bytes again is idempotent.
    """
    content = await file.read()
    filename = file.filename or "upload"
    suffix = Path(filename).suffix.lower()
    pages, metadata = _extract_pages(content, suffix)
    digest = sha1(content).hexdigest()[:16]
    return LoadedDocument(
        document_id=digest,
        source=source or filename,
        pages=pages,
        metadata={"filename": filename, "content_type": file.content_type or "", **metadata},
        paged=suffix == ".pdf",
    )


def load_file(path: Path, *, source: str | None = None) -> LoadedDocument:
    """Load a document from a local path (used by the seed script); mirrors load_upload."""
    content = path.read_bytes()
    suffix = path.suffix.lower()
    pages, metadata = _extract_pages(content, suffix)
    digest = sha1(content).hexdigest()[:16]
    return LoadedDocument(
        document_id=digest,
        source=source or path.name,
        pages=pages,
        metadata={"filename": path.name, **metadata},
        paged=suffix == ".pdf",
    )


def _extract_pages(content: bytes, suffix: str) -> tuple[list[str], dict]:
    """Extract plain text from raw bytes based on file type.

    PDFs keep one string per page (so citations can point at a page number); DOCX joins its
    paragraphs; everything else is decoded as UTF-8. Returns the pages plus type-specific
    metadata such as page_count or paragraph_count.
    """
    if suffix == ".pdf":
        reader = PdfReader(BytesIO(content))
        pages = [strip_page_numbers(page.extract_text() or "") for page in reader.pages]
        return pages, {"page_count": len(pages)}

    if suffix == ".docx":
        from docx import Document

        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        return ["\n".join(paragraphs)], {"paragraph_count": len(paragraphs)}

    return [content.decode("utf-8", errors="ignore")], {}


def strip_page_numbers(text: str) -> str:
    """Drop page-number lines, which would otherwise end up inside chunks and quoted passages.

    "Page N" always goes. "N of M" and "N / M" go only when they read as a page count
    (1 <= N <= M), so content such as "24/7" stays. A bare number stays too.
    """

    def page_number(match: re.Match[str]) -> str:
        labelled, current, total = match.group(1), int(match.group(2)), match.group(3)
        if total is None:
            return "" if labelled else match.group(0)
        return "" if 1 <= current <= int(total) else match.group(0)

    return _PAGE_NUMBER_LINE.sub(page_number, text).strip("\n")
