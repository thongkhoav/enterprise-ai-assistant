"""Build the bundled PDF and DOCX sample documents from the text below.

The generated files are committed, so this only needs to run after editing the text:

    uv run --with reportlab python scripts/build_seed_files.py

PDFs get one section group per page so answers can cite page numbers. The parental leave
policy goes to docs/demo_upload/ and is not seeded: it is the file an admin adds when an
employee asks about parental leave ("I don't know" before, a cited answer after).
"""

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parents[1]

# Each document: title, subtitle, then pages; each page is a list of (heading, paragraphs).
MSA = {
    "title": "Acme Master Service Agreement",
    "subtitle": "Version 4.2, effective January 1, 2026",
    "pages": [
        [
            (
                "1. Scope",
                [
                    "This Master Service Agreement (the Agreement) governs access to the Acme "
                    "platform by the customer named on the Order Form. It applies to the Starter, "
                    "Premium and Strategic plans. Where an Order Form and this Agreement conflict, "
                    "the Order Form controls for that order only.",
                    "Acme provides the subscribed services, standard support for the purchased "
                    "plan, and the product documentation published in the Acme help center. "
                    "Professional services such as custom integrations are ordered separately "
                    "under a statement of work.",
                ],
            ),
            (
                "2. Term and Renewal",
                [
                    "The initial subscription term is 12 months from the start date on the Order "
                    "Form, unless the Order Form states a multi-year term. Strategic plans may be "
                    "signed for 24 or 36 months.",
                    "Subscriptions renew automatically for successive 12-month terms unless either "
                    "party gives written notice of non-renewal at least 60 days before the end of "
                    "the current term. Notice sent by email to the legal contact on the Order Form "
                    "is sufficient.",
                    "Acme may increase fees at renewal by no more than 7 percent over the prior "
                    "term, and must give at least 90 days written notice of any increase.",
                ],
            ),
        ],
        [
            (
                "3. Fees and Payment",
                [
                    "Subscription fees are invoiced annually in advance. Invoices are payable "
                    "within 30 days of the invoice date (net 30). Fees are stated in USD and "
                    "exclude taxes, which the customer pays unless it provides a valid exemption "
                    "certificate.",
                    "Overdue amounts accrue interest of 1.5 percent per month or the maximum rate "
                    "allowed by law, whichever is lower. If an invoice is more than 30 days "
                    "overdue, Acme may suspend the services after giving 15 days written notice. "
                    "Suspension does not change the customer's payment obligations.",
                    "The customer must dispute an invoice in good faith within 20 days of "
                    "receiving it. Undisputed portions remain payable on the original due date. "
                    "Usage above the plan allowance is billed monthly in arrears at the overage "
                    "rates in the current price list.",
                ],
            ),
        ],
        [
            (
                "4. Service Level Agreement",
                [
                    "Acme commits to monthly uptime of 99.9 percent for the Premium plan and "
                    "99.95 percent for the Strategic plan. The Starter plan has no uptime "
                    "commitment. Uptime is measured per calendar month by Acme's external "
                    "monitoring of the production API and web application.",
                    "Scheduled maintenance is excluded from downtime. Maintenance runs on Sundays "
                    "between 02:00 and 04:00 UTC and is announced at least 72 hours in advance on "
                    "the Acme status page. Downtime caused by the customer, force majeure or "
                    "third-party networks outside Acme's control is also excluded.",
                    "If monthly uptime falls below the commitment, the customer receives a service "
                    "credit on its next invoice according to these tiers.",
                ],
            ),
            (
                "",
                [
                    "Uptime from 99.0 percent up to the commitment earns a credit of 10 percent "
                    "of the monthly fee.",
                    "Uptime from 95.0 percent up to 99.0 percent earns a credit of 25 percent of "
                    "the monthly fee.",
                    "Uptime below 95.0 percent earns a credit of 50 percent of the monthly fee.",
                ],
            ),
            (
                "",
                [
                    "Service credits are capped at 50 percent of the monthly fee for the affected "
                    "month. The customer must request a credit within 30 days after the end of the "
                    "affected month. Service credits are the customer's sole remedy for missing "
                    "the uptime commitment.",
                ],
            ),
        ],
        [
            (
                "5. Data Protection",
                [
                    "The customer owns all customer data. Acme processes customer data only to "
                    "provide the services and as described in the Data Processing Addendum.",
                    "Customer data is hosted in the region selected on the Order Form: United "
                    "States (AWS us-east-1) or European Union (AWS eu-central-1). Acme does not "
                    "move customer data out of the selected region except at the customer's "
                    "written request.",
                    "Acme notifies the customer of a confirmed security breach affecting customer "
                    "data within 48 hours of confirming it. Acme gives at least 30 days notice "
                    "before adding a new subprocessor, and the customer may object in writing "
                    "during that period.",
                    "After the Agreement ends, the customer may export its data for 30 days. Acme "
                    "deletes remaining customer data within 30 days after the export window "
                    "closes, except where law requires longer retention.",
                ],
            ),
        ],
        [
            (
                "6. Termination",
                [
                    "Either party may terminate the Agreement for material breach if the breach is "
                    "not cured within 30 days of written notice. If the customer terminates "
                    "because of Acme's uncured breach, Acme refunds prepaid fees for the remaining "
                    "term on a pro rata basis. Otherwise, fees paid are non-refundable.",
                ],
            ),
            (
                "7. Limitation of Liability",
                [
                    "Each party's total liability under the Agreement is limited to the fees paid "
                    "or payable by the customer in the 12 months before the event giving rise to "
                    "the claim. For breaches of the data protection section, the cap is two times "
                    "those fees. Neither party is liable for indirect or consequential damages, "
                    "including lost profits.",
                    "The caps do not apply to a party's indemnification obligations, to breach of "
                    "confidentiality, or to the customer's obligation to pay fees.",
                ],
            ),
            (
                "8. Governing Law",
                [
                    "The Agreement is governed by the laws of the State of Delaware, without "
                    "regard to conflict of laws rules. Disputes are resolved in the state or "
                    "federal courts located in Wilmington, Delaware.",
                ],
            ),
        ],
    ],
}

HANDBOOK = {
    "title": "Acme Employee Handbook",
    "subtitle": "2026 edition, People Operations",
    "pages": [
        [
            (
                "Welcome",
                [
                    "This handbook explains how we work at Acme. It applies to all full-time "
                    "employees. Country addenda add local rules where the law requires them, and "
                    "the addendum wins if it differs from this handbook.",
                ],
            ),
            (
                "Working Hours and Location",
                [
                    "Acme core hours are 10:00 to 15:00 in your local time zone. Outside core "
                    "hours you choose your schedule, as long as your team agrees on response "
                    "times for urgent work.",
                    "Office-based employees work from the office on Tuesdays and Thursdays and "
                    "may work remotely on the other days. Fully remote roles must be approved by "
                    "the department VP and are listed as remote in the job description.",
                    "Working from another country for more than 14 days in a calendar year needs "
                    "approval from People Operations, because it can create tax obligations.",
                ],
            ),
        ],
        [
            (
                "Paid Time Off",
                [
                    "Full-time employees receive 20 days of paid time off (PTO) per year. PTO "
                    "increases by one day for each completed year of service, up to 25 days. "
                    "Public holidays follow the calendar of the country you are employed in and "
                    "do not count against PTO.",
                    "Request PTO in Workday. Absences of three days or more should be requested "
                    "at least two weeks in advance. Up to 5 unused PTO days carry over to the next "
                    "year and expire on March 31.",
                ],
            ),
            (
                "Sick Days",
                [
                    "Employees receive 10 paid sick days per year. Tell your manager as early as "
                    "possible on the first day you are sick. A doctor's note is required after "
                    "three consecutive sick days.",
                    "Other leave types, such as bereavement or jury duty, are described in the "
                    "country addenda.",
                ],
            ),
        ],
        [
            (
                "Expenses",
                [
                    "Submit expenses in Expensify within 30 days of the purchase. Attach a receipt "
                    "for every expense above 25 USD. Expenses submitted after 60 days are not "
                    "reimbursed without VP approval.",
                    "Meals while travelling are reimbursed up to 75 USD per day for domestic "
                    "trips and 100 USD per day for international trips. Alcohol is not "
                    "reimbursed unless it is part of an approved customer dinner.",
                ],
            ),
            (
                "Business Travel",
                [
                    "All business trips need manager approval before booking. Book flights and "
                    "hotels through Navan at least 14 days before departure when possible.",
                    "Book economy class for flights under 6 hours and premium economy for flights "
                    "of 6 hours or longer. Hotels are capped at 250 USD per night, or 400 USD per "
                    "night in New York, London, San Francisco and Tokyo.",
                ],
            ),
        ],
        [
            (
                "Equipment",
                [
                    "Every employee receives a MacBook Pro, which is refreshed every three years. "
                    "Engineers may request an additional external monitor through the IT portal.",
                    "Remote and hybrid employees receive a one-time home office stipend of 500 USD "
                    "and an internet allowance of 40 USD per month, paid through payroll.",
                ],
            ),
            (
                "Security Basics",
                [
                    "Report a lost or stolen device to IT within one hour through the IT "
                    "portal, so the device can be locked and wiped. Never store customer data on "
                    "personal devices or personal cloud accounts.",
                    "All employees complete security awareness training when they join and every "
                    "year after that. Report suspected phishing with the Report Phishing button in "
                    "Gmail instead of clicking links in the email.",
                ],
            ),
        ],
        [
            (
                "Learning and Growth",
                [
                    "Each employee has a learning budget of 1,500 USD per year for courses, books "
                    "and certifications. Conference trips are paid from the same budget, with "
                    "travel booked under the travel policy.",
                    "Performance reviews happen twice a year, in April and October. Each review "
                    "includes a self-assessment, peer feedback and a conversation with your "
                    "manager about goals for the next six months.",
                ],
            ),
            (
                "Code of Conduct",
                [
                    "Treat colleagues, customers and partners with respect. Harassment and "
                    "discrimination are not tolerated. Report concerns to your manager, to People "
                    "Operations, or anonymously through the ethics hotline linked in the HR "
                    "portal. Acme does not allow retaliation against anyone who "
                    "reports a concern in good faith.",
                ],
            ),
        ],
    ],
}

BENEFITS = {
    "title": "Acme Benefits Guide 2026",
    "subtitle": "United States employees, People Operations",
    "pages": [
        [
            (
                "Who This Guide Covers",
                [
                    "This guide describes the 2026 benefits for full-time Acme employees in the "
                    "United States. Benefits start on your first day of employment. Employees in "
                    "other countries receive local benefits described in their country addendum.",
                ],
            ),
            (
                "Medical Plans",
                [
                    "Acme offers two medical plans through Aetna: the Acme PPO and the Acme HDHP "
                    "with a health savings account (HSA). Both plans use the same national "
                    "network and cover preventive care at 100 percent.",
                    "Acme pays 90 percent of the monthly premium for employees and 75 percent for "
                    "spouses, domestic partners and children. The PPO has a 500 USD individual "
                    "deductible; the HDHP has a 1,700 USD individual deductible.",
                ],
            ),
            (
                "Dental and Vision",
                [
                    "Dental coverage through Delta Dental and vision coverage through VSP are "
                    "free for employees. Acme pays 50 percent of the dental and vision premium "
                    "for dependents. Vision covers one eye exam and up to 200 USD for glasses or "
                    "contact lenses every 12 months.",
                ],
            ),
        ],
        [
            (
                "401(k) Retirement Plan",
                [
                    "Acme matches 100 percent of your 401(k) contributions up to 4 percent of "
                    "your base salary. The match is paid every payroll and vests immediately. "
                    "You can enroll, change your contribution or stop at any time in Fidelity "
                    "NetBenefits; changes apply from the next payroll.",
                    "New employees are enrolled automatically at a 3 percent contribution after "
                    "30 days unless they choose a different rate or opt out.",
                ],
            ),
            (
                "Health Savings Account",
                [
                    "Employees on the Acme HDHP receive an Acme HSA contribution of 750 USD per "
                    "year for individual coverage or 1,500 USD per year for family coverage, "
                    "paid in equal amounts each payroll. The HSA belongs to you and stays with "
                    "you if you leave Acme.",
                ],
            ),
            (
                "Life and Disability Insurance",
                [
                    "Acme provides life insurance of two times your base salary, up to 500,000 "
                    "USD, at no cost. Short-term disability pays 70 percent of base salary for "
                    "up to 12 weeks after a 7-day waiting period. Long-term disability pays 60 "
                    "percent of base salary after 90 days.",
                ],
            ),
        ],
        [
            (
                "Wellbeing",
                [
                    "Every employee receives a wellness stipend of 75 USD per month, paid through "
                    "payroll, for gym memberships, fitness classes, meditation apps or sports "
                    "equipment. No receipts are needed.",
                    "The Employee Assistance Program (EAP) through Lyra offers 8 free counseling "
                    "sessions per person per year, for employees and their family members. Lyra "
                    "is available 24/7 and sessions are confidential; Acme never sees who uses it.",
                    "The commuter benefit lets office-based employees pay for transit and parking "
                    "with up to 315 USD per month of pre-tax salary.",
                ],
            ),
            (
                "Enrolling and Making Changes",
                [
                    "Choose your benefits in Workday within 30 days of your start date. If you "
                    "do not choose, you are enrolled in the Acme PPO for yourself only.",
                    "Open enrollment for 2027 runs from November 3 to November 17, 2026. Outside "
                    "open enrollment you can change your plans only within 30 days of a "
                    "qualifying life event, such as marriage, a new dependent or loss of other "
                    "coverage.",
                    "Questions about benefits go to People Operations through the HR portal.",
                ],
            ),
        ],
    ],
}

IT_HELP_DESK = {
    "title": "Acme IT Help Desk Guide",
    "sections": [
        (
            "Getting Help",
            [
                "Open a ticket in the IT portal or ask in the #it-help Slack channel. The help "
                "desk is staffed Monday to Friday from 08:00 to 20:00 in the US Eastern and "
                "Central European time zones.",
                "Response targets depend on priority. P1, when you cannot work at all, gets a "
                "response within 30 minutes. P2, when something important is broken but you can "
                "still work, gets a response within 4 business hours. P3 requests and questions "
                "get a response within 2 business days.",
                "Outside help desk hours, use the Urgent option in the IT portal for P1 issues. "
                "It pages the on-call IT engineer.",
            ],
        ),
        (
            "Passwords and Multi-Factor Authentication",
            [
                "Acme uses Okta for single sign-on. Passwords must be at least 14 characters "
                "and are changed every 12 months, or immediately if you suspect someone knows "
                "yours.",
                "To reset a forgotten password, click Need help signing in on the Okta page and "
                "follow the self-service steps with Okta Verify. If you lost the phone with Okta "
                "Verify, the help desk verifies your identity on a video call before resetting "
                "your factors.",
                "Multi-factor authentication is required for every account. Use Okta Verify "
                "push or a YubiKey; SMS codes are not allowed.",
            ],
        ),
        (
            "VPN",
            [
                "The GlobalProtect VPN is required for internal tools such as the admin console, "
                "Grafana and the staging environment. It is not needed for Google Workspace, "
                "Slack, Zoom or Workday.",
                "Install GlobalProtect from the Self Service app on your Mac and sign in with "
                "Okta. If the VPN will not connect, check that your Mac's date and time are set "
                "automatically, quit and reopen GlobalProtect, and switch to another network. If "
                "it still fails, open a P2 ticket with a screenshot of the error.",
            ],
        ),
        (
            "Laptops and Equipment",
            [
                "Laptops are refreshed every three years. The IT portal shows when your laptop "
                "is due and sends the new one five business days before the swap.",
                "If your laptop is damaged, open a P2 ticket. Office-based employees get a "
                "loaner laptop the same day; remote employees receive a replacement within two "
                "business days.",
                "Report a lost or stolen laptop or phone within one hour as a P1 ticket, so IT "
                "can lock and wipe it.",
            ],
        ),
        (
            "Software Requests",
            [
                "Apps in the Self Service catalog are pre-approved and install immediately. To "
                "request anything else, open a software request in the IT portal. It needs "
                "manager approval and a security review, which takes up to five business days.",
                "Do not buy software licences with a corporate card or expense them without IT "
                "approval, even for free trials that ask for a card.",
            ],
        ),
        (
            "Office Wi-Fi",
            [
                "Acme offices have two Wi-Fi networks. Employees connect laptops and phones to "
                "the staff network, Acme-Staff, and sign in with Okta. Acme-Staff reaches office "
                "printers and internal tools without the VPN.",
                "Visitors, and personal devices that are not enrolled with IT, use the guest "
                "Wi-Fi network, Acme-Guest. Reception gives out the guest Wi-Fi password, which "
                "changes every day. The guest network is internet-only and cannot reach "
                "printers or internal tools.",
            ],
        ),
        (
            "Printing",
            [
                "Printers use badge release: send your document to Acme-Print, then tap your "
                "badge on any office printer to print it. Print jobs are deleted after 24 "
                "hours if nobody releases them.",
            ],
        ),
    ],
}

PARENTAL_LEAVE = {
    "title": "Acme Parental Leave Policy",
    "subtitle": "Effective January 1, 2026, People Operations",
    "pages": [
        [
            (
                "Purpose",
                [
                    "Acme wants every parent to have time to bond with a new child without "
                    "worrying about pay. This policy applies to all full-time employees and "
                    "replaces the parental leave section of earlier handbooks.",
                ],
            ),
            (
                "Eligibility",
                [
                    "Employees are eligible after 90 days of employment. The policy covers the "
                    "birth of a child, adoption, foster placement and surrogacy, and applies to "
                    "every parent regardless of gender.",
                ],
            ),
            (
                "Paid Leave",
                [
                    "Eligible employees receive 16 weeks of parental leave at 100 percent of base "
                    "salary. Leave can be taken in one block or in up to two blocks, and must be "
                    "used within 12 months of the birth or placement.",
                    "Birthing parents may start their leave up to 2 weeks before the due date, "
                    "and those weeks count toward the 16 weeks. Where local law gives more leave "
                    "or higher pay, the local law applies.",
                ],
            ),
        ],
        [
            (
                "Requesting Leave",
                [
                    "Tell your manager and submit a parental leave request in Workday at least 8 "
                    "weeks before the planned start date. For adoption or foster placement, "
                    "submit the request as soon as you know the date. People Operations confirms "
                    "the request within 3 business days.",
                ],
            ),
            (
                "During Your Leave",
                [
                    "Health, dental and vision coverage, the 401(k) match and the wellness "
                    "stipend continue during leave. PTO keeps accruing as usual.",
                ],
            ),
            (
                "Returning to Work",
                [
                    "Your role and pay are protected while you are on leave. After returning, you "
                    "may work a 60 percent schedule at full pay for up to 4 weeks to ease back "
                    "in. Agree on the schedule with your manager before your leave ends.",
                ],
            ),
        ],
    ],
}


def build_pdf(document: dict, path: Path) -> None:
    styles = getSampleStyleSheet()
    story = [
        Paragraph(document["title"], styles["Title"]),
        Paragraph(document["subtitle"], styles["Italic"]),
        Spacer(1, 0.2 * inch),
    ]
    for index, page in enumerate(document["pages"]):
        if index:
            story.append(PageBreak())
        for heading, body in page:
            if heading:
                story.append(Paragraph(heading, styles["Heading2"]))
            story.extend(Paragraph(paragraph, styles["BodyText"]) for paragraph in body)

    def footer(canvas, doc) -> None:
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(LETTER[0] - inch, 0.6 * inch, f"Page {doc.page}")

    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=LETTER, title=document["title"]).build(
        story, onFirstPage=footer, onLaterPages=footer
    )


def build_docx(document: dict, path: Path) -> None:
    docx = Document()
    docx.add_heading(document["title"], level=0)
    for heading, body in document["sections"]:
        docx.add_heading(heading, level=1)
        for paragraph in body:
            docx.add_paragraph(paragraph)
    docx.save(str(path))


if __name__ == "__main__":
    build_pdf(MSA, ROOT / "docs/seed/acme_master_service_agreement.pdf")
    build_pdf(HANDBOOK, ROOT / "docs/seed/acme_employee_handbook.pdf")
    build_pdf(BENEFITS, ROOT / "docs/seed/acme_benefits_guide_2026.pdf")
    build_docx(IT_HELP_DESK, ROOT / "docs/seed/acme_it_help_desk_guide.docx")
    build_pdf(PARENTAL_LEAVE, ROOT / "docs/demo_upload/acme_parental_leave_policy_2026.pdf")
    print("Built the seed PDFs, the IT help desk DOCX and the parental leave policy.")
